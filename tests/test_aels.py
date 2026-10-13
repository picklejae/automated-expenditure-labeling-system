"""Tests for the labeling pipeline and the web API. Run with: pytest"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from aels.app import create_app
from aels.db import EXCLUDED_SECTION, LabelStore
from aels.processing import (
    FIELDS,
    collapse_offsets,
    export_labeled,
    process_sof,
    read_sof,
    totals_by_section,
)

HEADERS = [
    "Fund", "Funds Center", "Area", "Commitment Item", "Fiscal Month",
    "Line Item Text", "Obligation Doc", "Posting Date", "Reference",
    "WBS Element", "Mission", "Comm. (Cum)", "Oblig. (Cum)",
]


def make_row(code, desc, doc, amount):
    return ["F", "FC", "A", code, 4, desc, doc, "2025-04-01", "R", "W", "M", amount, amount]


def write_xlsx(path, rows, header=True):
    wb = Workbook()
    ws = wb.active
    if header:
        ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return str(path)


@pytest.fixture
def store(tmp_path):
    return LabelStore(str(tmp_path / "test.db"))


# ---- reading ---------------------------------------------------------------

def test_reads_file_with_header(tmp_path):
    path = write_xlsx(tmp_path / "a.xlsx", [make_row("2610", "PAPER", "D1", 10.5)])
    items = read_sof(path)
    assert len(items) == 1
    assert items[0]["commitment_item"] == "2610"
    assert items[0]["amount"] == 10.5


def test_reads_file_without_header(tmp_path):
    path = write_xlsx(tmp_path / "a.xlsx", [make_row("2610", "PAPER", "D1", 10.5)], header=False)
    items = read_sof(path)
    assert len(items) == 1 and items[0]["description"] == "PAPER"


def test_header_columns_can_be_reordered(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Oblig. (Cum)", "Line Item Text", "Commitment Item", "Obligation Doc"])
    ws.append([99.0, "TONER", "2610", "D9"])
    wb.save(tmp_path / "a.xlsx")
    items = read_sof(str(tmp_path / "a.xlsx"))
    assert items[0]["amount"] == 99.0 and items[0]["description"] == "TONER"


def test_rejects_unrecognized_layout(tmp_path):
    wb = Workbook()
    wb.active.append(["just", "three", "cols"])
    wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(ValueError):
        read_sof(str(tmp_path / "bad.xlsx"))


def test_skips_blank_rows(tmp_path):
    rows = [make_row("2610", "PAPER", "D1", 10), [None] * len(FIELDS)]
    assert len(read_sof(write_xlsx(tmp_path / "a.xlsx", rows))) == 1


def test_hand_labeled_section_column_is_imported(tmp_path, store):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS + ["Section"])
    ws.append(make_row("2610", "PAPER", "D1", 10) + ["Supply"])
    ws.append(make_row("2650", "MONITOR", "D2", 400) + [None])
    wb.save(tmp_path / "labeled.xlsx")

    result = process_sof(str(tmp_path / "labeled.xlsx"), store)
    assert result["stats"]["imported_labels"] == 1
    assert store.get_labels() == {("2610", "PAPER"): "Supply"}
    assert result["stats"]["needs_label"] == 1


def test_section_column_without_header_is_imported(tmp_path, store):
    path = write_xlsx(tmp_path / "a.xlsx", [make_row("2610", "PAPER", "D1", 10) + ["HQ Co"]], header=False)
    process_sof(path, store)
    assert store.get_labels()[("2610", "PAPER")] == "HQ Co"


# ---- offsets ----------------------------------------------------------------

def _item(doc, desc, amount, row):
    return {"obligation_doc": doc, "description": desc, "commitment_item": "2610",
            "amount": amount, "row_index": row}


def test_full_reversal_is_dropped():
    items = [_item("D1", "TIRE", 250, 1), _item("D1", "TIRE", -250, 2)]
    assert collapse_offsets(items) == []


def test_partial_reversal_nets_to_one_line():
    items = [_item("D1", "TIRE", 200, 1), _item("D1", "TIRE", -125, 2)]
    result = collapse_offsets(items)
    assert len(result) == 1 and result[0]["amount"] == 75


def test_zero_dollar_line_is_dropped():
    assert collapse_offsets([_item("D1", "TIRE", 0, 1)]) == []


def test_same_description_on_different_docs_is_kept_separate():
    items = [_item("D1", "TIRE", 100, 1), _item("D2", "TIRE", 100, 2)]
    assert len(collapse_offsets(items)) == 2


def test_floating_point_noise_counts_as_zero():
    items = [_item("D1", "X", 0.1, 1), _item("D1", "X", 0.2, 2), _item("D1", "X", -0.3, 3)]
    assert collapse_offsets(items) == []


# ---- memory and change detection --------------------------------------------

def test_labels_are_remembered_across_reports(tmp_path, store):
    p1 = write_xlsx(tmp_path / "p1.xlsx", [make_row("2610", "PAPER", "D1", 10)])
    first = process_sof(p1, store)
    assert first["stats"]["needs_label"] == 1

    store.save_label("2610", "PAPER", "Supply")

    p2 = write_xlsx(tmp_path / "p2.xlsx", [
        make_row("2610", "PAPER", "D7", 30),   # same item, new obligation
        make_row("2650", "MONITOR", "D8", 400),
    ])
    second = process_sof(p2, store)
    by_desc = {i["description"]: i for i in second["items"]}
    assert by_desc["PAPER"]["section"] == "Supply"
    assert by_desc["PAPER"]["label_source"] == "auto"
    assert by_desc["MONITOR"]["section"] == ""


def test_relabeling_overwrites_previous_label(store):
    store.save_label("2610", "PAPER", "Supply")
    store.save_label("2610", "PAPER", "Admin")
    assert store.get_labels()[("2610", "PAPER")] == "Admin"


def test_changed_amounts_are_flagged(tmp_path, store):
    process_sof(write_xlsx(tmp_path / "p1.xlsx", [
        make_row("2625", "DIESEL", "D1", 1000),
        make_row("2610", "PAPER", "D2", 50),
    ]), store)
    result = process_sof(write_xlsx(tmp_path / "p2.xlsx", [
        make_row("2625", "DIESEL", "D1", 1400),
        make_row("2610", "PAPER", "D2", 50),
    ]), store)
    changed = [i for i in result["items"] if i["changed"]]
    assert len(changed) == 1
    assert changed[0]["previous_amount"] == 1000 and changed[0]["amount"] == 1400


def test_new_lines_are_marked(tmp_path, store):
    process_sof(write_xlsx(tmp_path / "p1.xlsx", [make_row("2610", "PAPER", "D1", 10)]), store)
    result = process_sof(write_xlsx(tmp_path / "p2.xlsx", [
        make_row("2610", "PAPER", "D1", 10),
        make_row("2650", "MONITOR", "D2", 400),
    ]), store)
    status = {i["description"]: i["status_text"] for i in result["items"]}
    assert status == {"PAPER": "", "MONITOR": "New"}


def test_charges_from_other_units_are_kept_out_of_unit_total():
    items = [
        {"section": "Supply", "amount": 100.0},
        {"section": EXCLUDED_SECTION, "amount": 2750.0},
        {"section": "", "amount": 5.0},
    ]
    totals = totals_by_section(items)
    assert totals["unit_total"] == 105.0
    assert totals["not_our_unit"] == 2750.0
    assert EXCLUDED_SECTION not in totals["sections"]


# ---- export -------------------------------------------------------------------

def test_export_has_labels_and_summary(tmp_path):
    items = [
        {"commitment_item": "2610", "description": "PAPER", "obligation_doc": "D1",
         "amount": 10.0, "section": "Supply", "label_source": "auto"},
        {"commitment_item": "2650", "description": "MONITOR", "obligation_doc": "D2",
         "amount": 400.0, "section": "", "label_source": "", "status_text": "New"},
        {"commitment_item": "2680", "description": "TRAILER", "obligation_doc": "D3",
         "amount": 50.0, "section": EXCLUDED_SECTION, "label_source": "manual"},
    ]
    out = tmp_path / "out.xlsx"
    export_labeled(items, str(out))
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Labeled"]
    assert wb["Labeled"]["H2"].value == "Supply"
    assert wb["Summary"]["A2"].value == "Supply"
    assert wb["Summary"]["B2"].value.startswith("=SUMIF(")
    assert wb["Labeled"]["I3"].value == "New"
    summary_labels = [c.value for c in wb["Summary"]["A"]]
    assert "Unit total" in summary_labels and "Report total" in summary_labels


# ---- API ------------------------------------------------------------------------

@pytest.fixture
def client(tmp_path):
    app = create_app(db_path=str(tmp_path / "api.db"), data_dir=str(tmp_path / "data"))
    return app.test_client()


def _upload(client, tmp_path, rows, name="report.xlsx"):
    path = write_xlsx(tmp_path / name, rows)
    with open(path, "rb") as f:
        data = {"file": (io.BytesIO(f.read()), name)}
    return client.post("/api/upload", data=data, content_type="multipart/form-data")


def test_api_upload_label_export_flow(client, tmp_path):
    res = _upload(client, tmp_path, [make_row("2610", "PAPER", "D1", 10)])
    assert res.status_code == 200
    body = res.get_json()
    sid = body["session_id"]
    assert body["stats"]["needs_label"] == 1

    res = client.post("/api/label", json={"session_id": sid, "commitment_item": "2610",
                                          "description": "PAPER", "section": "Supply"})
    totals = res.get_json()["summary"]["totals"]
    assert totals["sections"] == {"Supply": 10.0} and totals["unit_total"] == 10.0

    res = client.post("/api/export", json={"session_id": sid})
    assert res.status_code == 200
    assert res.headers["Content-Disposition"].endswith("report_LABELED.xlsx")

    # The label is remembered on the next upload
    res = _upload(client, tmp_path, [make_row("2610", "PAPER", "D5", 20)], name="next.xlsx")
    assert res.get_json()["stats"]["auto_labeled"] == 1


def test_api_rejects_non_excel(client):
    res = client.post("/api/upload", data={"file": (io.BytesIO(b"x"), "notes.txt")},
                      content_type="multipart/form-data")
    assert res.status_code == 400


def test_api_unknown_session(client):
    res = client.post("/api/export", json={"session_id": "nope"})
    assert res.status_code == 404


def test_api_add_section(client):
    res = client.post("/api/sections", json={"name": "Security"})
    assert "Security" in res.get_json()["sections"]
