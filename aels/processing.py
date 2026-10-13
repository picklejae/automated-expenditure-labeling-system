"""Reading, cleaning, labeling and exporting Status of Funds (SOF) reports.

Pipeline for one uploaded report:

    read_sof()          Excel -> list of line items (works with or without a header row)
    collapse_offsets()  net +/- entries on the same obligation; drop lines that net to $0
    apply_labels()      fill in sections the database already knows
    flag_changes()      compare with the previous report: mark lines New or Changed
    export_labeled()    write a labeled workbook with a totals-by-section summary
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import EXCLUDED_SECTION, LabelStore

# Fields the tool uses, in the column order of the raw report export.
# When the file has a header row, columns are matched by name instead,
# so reordered or extra columns still work.
FIELDS = [
    "fund",
    "funds_center",
    "area",
    "commitment_item",
    "fiscal_month",
    "description",
    "obligation_doc",
    "posting_date",
    "reference",
    "wbs_element",
    "mission",
    "commitment_amount",
    "amount",  # Oblig. (Cum) -- the dollar value that matters
]

HEADER_ALIASES = {
    "fund": ["fund"],
    "funds_center": ["funds center", "fund center"],
    "area": ["area", "functional area"],
    "commitment_item": ["commitment item", "cmmt item"],
    "fiscal_month": ["fiscal month", "period"],
    "description": ["line item text", "description"],
    "obligation_doc": ["obligation doc", "obligation document", "document"],
    "posting_date": ["posting date"],
    "reference": ["reference"],
    "wbs_element": ["wbs element", "wbs"],
    "mission": ["mission"],
    "commitment_amount": ["comm. (cum)", "commitment (cum)"],
    "amount": ["oblig. (cum)", "obligation (cum)", "obligated"],
    # Optional: a section column the user filled in by hand in an earlier
    # report. Its labels are imported into the database.
    "file_section": ["section", "label"],
}

EXPORT_HEADERS = [
    ("Commitment Item", "commitment_item"),
    ("Line Item Text", "description"),
    ("Obligation Doc", "obligation_doc"),
    ("Posting Date", "posting_date"),
    ("Reference", "reference"),
    ("Mission", "mission"),
    ("Oblig. (Cum)", "amount"),
    ("Section", "section"),
    ("Status", "status_text"),
]

CENT = 0.005  # amounts closer than half a cent are treated as equal


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------

def _clean(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _to_float(value: Any) -> float:
    try:
        number = float(str(value).replace(",", "").replace("$", ""))
    except (TypeError, ValueError):
        return 0.0
    return 0.0 if pd.isna(number) else number


def _map_header(row: list[Any]) -> dict[str, int] | None:
    """If `row` looks like a header row, return {field: column index}."""
    names = [_clean(v).lower() for v in row]
    mapping = {}
    for field, aliases in HEADER_ALIASES.items():
        for i, name in enumerate(names):
            if name in aliases:
                mapping[field] = i
                break
    required = {"commitment_item", "description", "obligation_doc", "amount"}
    return mapping if required <= mapping.keys() else None


def read_sof(filepath: str) -> list[dict[str, Any]]:
    """Read the first sheet of a SOF report into a list of line-item dicts."""
    df = pd.read_excel(filepath, header=None, dtype=object)
    if df.empty:
        return []

    columns = _map_header(df.iloc[0].tolist())
    start = 1
    if columns is None:  # no header row: fall back to the standard column order
        if df.shape[1] < len(FIELDS):
            raise ValueError(
                f"Expected a header row or at least {len(FIELDS)} columns, "
                f"found {df.shape[1]}."
            )
        columns = {field: i for i, field in enumerate(FIELDS)}
        if df.shape[1] > len(FIELDS):  # an extra column after the amounts holds sections
            columns["file_section"] = len(FIELDS)
        start = 0

    items = []
    for row_index in range(start, len(df)):
        row = df.iloc[row_index]
        item = {field: _clean(row.iloc[i]) for field, i in columns.items()}
        if not item["commitment_item"]:
            continue  # blank or subtotal row
        item["amount"] = _to_float(row.iloc[columns["amount"]])
        item.setdefault("file_section", "")
        item["row_index"] = row_index
        items.append(item)
    return items


# --------------------------------------------------------------------------
# Cleaning
# --------------------------------------------------------------------------

def collapse_offsets(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Net entries that share an obligation doc + description.

    Reports often contain a charge and its reversal as two lines
    (e.g. +$120 and -$120). Those are netted into one line, and any line
    that nets to $0 is dropped.
    """
    groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        groups[(item["obligation_doc"], item["description"])].append(item)

    collapsed = []
    for group in groups.values():
        net = round(sum(i["amount"] for i in group), 2)
        if abs(net) < CENT:
            continue
        latest = dict(max(group, key=lambda i: i["row_index"]))
        latest["amount"] = net
        if not latest.get("file_section"):
            latest["file_section"] = next((i["file_section"] for i in group if i.get("file_section")), "")
        collapsed.append(latest)
    collapsed.sort(key=lambda i: i["row_index"])
    return collapsed


# --------------------------------------------------------------------------
# Labeling and change detection
# --------------------------------------------------------------------------

def label_key(item: dict[str, Any]) -> tuple[str, str]:
    return (item["commitment_item"], item["description"])


def history_key(item: dict[str, Any]) -> tuple[str, str, str]:
    return (item["obligation_doc"], item["commitment_item"], item["description"])


def apply_labels(items: list[dict[str, Any]], known: dict[tuple[str, str], str]) -> None:
    for item in items:
        section = known.get(label_key(item))
        item["section"] = section or ""
        item["label_source"] = "auto" if section else ""


def flag_changes(
    items: list[dict[str, Any]], history: dict[tuple[str, str, str], float]
) -> None:
    for item in items:
        previous = history.get(history_key(item))
        item["previous_amount"] = previous
        item["is_new"] = previous is None
        item["changed"] = previous is not None and abs(previous - item["amount"]) >= CENT
        if item["is_new"]:
            item["status_text"] = "New"
        elif item["changed"]:
            item["status_text"] = f"Changed (was {previous:,.2f})"
        else:
            item["status_text"] = ""


def process_sof(filepath: str, store: LabelStore) -> dict[str, Any]:
    """Run the full pipeline on one report and record its amounts."""
    raw = read_sof(filepath)
    items = collapse_offsets(raw)
    for i, item in enumerate(items):
        item["id"] = i

    # Labels typed into a Section column of the file are imported first, so a
    # report that was labeled by hand in Excel seeds the database.
    imported = {label_key(i): i["file_section"] for i in items if i["file_section"]}
    if imported:
        store.save_labels((ci, desc, sec) for (ci, desc), sec in imported.items())
        for section in set(imported.values()):
            store.add_section(section)

    apply_labels(items, store.get_labels())
    flag_changes(items, store.get_history())
    store.save_history(
        (i["obligation_doc"], i["commitment_item"], i["description"], i["amount"])
        for i in items
    )

    return {
        "items": items,
        "raw_count": len(raw),
        "stats": {
            "total": len(items),
            "auto_labeled": sum(1 for i in items if i["section"]),
            "needs_label": sum(1 for i in items if not i["section"]),
            "changed": sum(1 for i in items if i["changed"]),
            "new": sum(1 for i in items if i["is_new"]),
            "offsets_removed": len(raw) - len(items),
            "imported_labels": len(imported),
        },
    }


def assign_label(items: list[dict[str, Any]], commitment_item: str, description: str,
                 section: str) -> int:
    """Apply a label to every matching line in the current report. Returns count."""
    count = 0
    for item in items:
        if label_key(item) == (commitment_item, description):
            item["section"] = section
            item["label_source"] = "manual"
            count += 1
    return count


def totals_by_section(items: list[dict[str, Any]]) -> dict[str, Any]:
    """Totals per section for the unit, with erroneous charges kept separate."""
    totals: dict[str, float] = defaultdict(float)
    excluded = 0.0
    for item in items:
        if item["section"] == EXCLUDED_SECTION:
            excluded += item["amount"]
        else:
            totals[item["section"] or "(unlabeled)"] += item["amount"]
    sections = {k: round(v, 2) for k, v in sorted(totals.items())}
    return {
        "sections": sections,
        "unit_total": round(sum(sections.values()), 2),
        "not_our_unit": round(excluded, 2),
    }


# --------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FILLS = {
    "auto": PatternFill("solid", start_color="E6F4EA"),     # green: from memory
    "manual": PatternFill("solid", start_color="FFF6D5"),   # yellow: labeled this session
    "": PatternFill("solid", start_color="FDE7E7"),         # red: still unlabeled
}
STATUS_FILLS = {
    "New": PatternFill("solid", start_color="E3EEFF"),
    "Changed": PatternFill("solid", start_color="FFE9D2"),
}
MONEY = '#,##0.00;[Red]-#,##0.00'


def _write_header(ws, headers: list[str]) -> None:
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def export_labeled(items: list[dict[str, Any]], out_path: str) -> None:
    """Write the labeled lines plus a Summary sheet that totals by section.

    Summary totals are SUMIF formulas, so they stay correct if someone
    edits a section in Excel afterwards. Lines labeled "Not our unit" are
    totaled separately so they can be disputed and are not counted as the
    unit's spending.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Labeled"
    _write_header(ws, [h for h, _ in EXPORT_HEADERS])

    for r, item in enumerate(items, 2):
        for c, (_, field) in enumerate(EXPORT_HEADERS, 1):
            value = item.get(field)
            ws.cell(row=r, column=c, value=None if value == "" else value)
        ws.cell(row=r, column=7).number_format = MONEY
        ws.cell(row=r, column=8).fill = FILLS.get(item.get("label_source", ""), FILLS[""])
        status = (item.get("status_text") or "").split(" ")[0]
        if status in STATUS_FILLS:
            ws.cell(row=r, column=9).fill = STATUS_FILLS[status]

    widths = [16, 38, 24, 13, 16, 28, 15, 14, 24]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.auto_filter.ref = f"A1:I{max(len(items) + 1, 1)}"

    summary = wb.create_sheet("Summary")
    _write_header(summary, ["Section", "Total Oblig. (Cum)", "Lines"])
    last = len(items) + 1
    sec_rng = f"Labeled!$H$2:$H${last}"
    amt_rng = f"Labeled!$G$2:$G${last}"
    bold = Font(bold=True)

    def write_row(r, name, total_formula, count_formula, font=None):
        summary.cell(row=r, column=1, value=name)
        summary.cell(row=r, column=2, value=total_formula).number_format = MONEY
        summary.cell(row=r, column=3, value=count_formula)
        if font:
            for c in (1, 2, 3):
                summary.cell(row=r, column=c).font = font

    sections = sorted({i["section"] for i in items if i["section"] and i["section"] != EXCLUDED_SECTION})
    r = 2
    for section in sections:
        write_row(r, section, f"=SUMIF({sec_rng},A{r},{amt_rng})", f"=COUNTIF({sec_rng},A{r})")
        r += 1
    write_row(r, "(unlabeled)", f'=SUMIF({sec_rng},"",{amt_rng})', f"=COUNTBLANK({sec_rng})")
    unit_row = r + 1
    write_row(unit_row, "Unit total", f"=SUM(B2:B{r})", f"=SUM(C2:C{r})", bold)

    excl_row = unit_row + 2
    write_row(excl_row, f"{EXCLUDED_SECTION} (charged in error)",
              f'=SUMIF({sec_rng},"{EXCLUDED_SECTION}",{amt_rng})',
              f'=COUNTIF({sec_rng},"{EXCLUDED_SECTION}")')
    write_row(excl_row + 1, "Report total", f"=B{unit_row}+B{excl_row}", f"=C{unit_row}+C{excl_row}", bold)

    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 20
    summary.column_dimensions["C"].width = 10

    wb.move_sheet("Summary", offset=-1)
    wb.active = 0
    wb.save(out_path)
