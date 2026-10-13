"""Generate two fake, consecutive Status of Funds reports for demos and tests.

Everything in these files is invented: units, documents, people, amounts.
The goal is a realistic shape, not real data:

  * one fiscal year of obligations (Oct-Sep) totaling roughly $1M
  * 1,000+ lines per report after offsets are netted out
  * travel (airfare, lodging, per diem, rental cars), oddly named training
    courses, card purchases, supply requisitions, repair parts, service
    contracts with monthly funding increments, fuel, medical and IT
  * returns, cancellations and voucher settlements that post as offsetting
    lines, and a handful of charges that belong to other units

Run from the repo root:

    python scripts/make_sample_data.py

Writes sample_data/SOF_period1.xlsx (as of 15 Aug, already labeled by hand in
a Section column, the way the analyst would have done it last cycle) and
sample_data/SOF_period2.xlsx (as of 29 Aug, a fresh unlabeled export).
"""

from __future__ import annotations

import random
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

FY_START = date(2024, 10, 1)
PERIOD1_AS_OF = date(2025, 8, 15)
PERIOD2_AS_OF = date(2025, 8, 29)
FY_END = date(2025, 9, 30)

HEADERS = [
    "Fund", "Funds Center", "Area", "Commitment Item", "Fiscal Month",
    "Line Item Text", "Obligation Doc", "Posting Date", "Reference",
    "WBS Element", "Mission", "Comm. (Cum)", "Oblig. (Cum)",
]

OUR_FC = "FC-1001"
COMPANIES = ["HQ Co", "Alpha Co", "Bravo Co", "Charlie Co"]
NOT_OURS = "Not our unit"
SCALE = 0.84  # tunes the fiscal-year total to roughly $1M


@dataclass
class Row:
    code: str
    desc: str
    doc: str
    posted: date
    amount: float
    mission: str
    reference: str
    section: str
    funds_center: str = OUR_FC


rng = random.Random(2025)
rows: list[Row] = []
_counters: dict[str, int] = {}


def next_id(prefix: str, width: int = 5) -> str:
    _counters[prefix] = _counters.get(prefix, 0) + 1
    return f"{prefix}{_counters[prefix]:0{width}d}"


def rand_date(start: date = FY_START, end: date = FY_END) -> date:
    return start + timedelta(days=rng.randint(0, (end - start).days))


def money(low: float, high: float) -> float:
    return round(rng.uniform(low, high), 2)


def add(code, desc, doc, posted, amount, mission, reference, section, fc=OUR_FC,
        exact=False) -> Row:
    if not exact:
        amount *= SCALE
    row = Row(code, desc, doc, posted, round(amount, 2), mission, reference, section, fc)
    rows.append(row)
    return row


def add_adjustment(row: Row, amount: float, days_later: tuple[int, int] = (5, 40)) -> None:
    """A later posting against the same obligation (return, refund, settlement)."""
    when = row.posted + timedelta(days=rng.randint(*days_later))
    add(row.code, row.desc, row.doc, min(when, FY_END), amount, row.mission,
        row.reference, row.section, row.funds_center, exact=True)


# ---------------------------------------------------------------------------
# Travel: each trip is one travel authorization with several lines.
# ---------------------------------------------------------------------------

TRIP_PURPOSES = [
    "ADV LDR CRS", "BN S4 CONF", "SITE SURVEY", "MAINT MGMT CONF", "PRE-CMD CRS",
    "MEDIC RECERT", "RANGE RECON", "LOG SYNC", "FIELD TRIAL SPT", "EXERCISE PLNG CONF",
    "PROPERTY BOOK TRNG", "UNIT MVMT OFFICER CRS", "SAFETY STANDDOWN SPT", "AWARDS BOARD",
    "TRANSITION BRIEF", "ARMS ROOM INSP", "COMSEC AUDIT", "MASTER DRIVER CRS",
]
TRIP_LOCATIONS = [
    ("FT MOORE GA", 120, 68), ("FT LIBERTY NC", 110, 64), ("JBLM WA", 155, 74),
    ("FT LEONARD WOOD MO", 98, 59), ("ABERDEEN PG MD", 145, 74), ("HUNTSVILLE AL", 118, 64),
    ("FT BLISS TX", 104, 64), ("SAN ANTONIO TX", 139, 69), ("FT CARSON CO", 129, 69),
    ("FT CAMPBELL KY", 107, 64), ("FT JOHNSON LA", 98, 59), ("ARLINGTON VA", 258, 79),
    ("FT IRWIN CA", 110, 64), ("FT GREGG-ADAMS VA", 112, 64), ("CAMP ATTERBURY IN", 104, 59),
]


def make_trips(count: int) -> None:
    for _ in range(count):
        purpose = rng.choice(TRIP_PURPOSES)
        loc, lodging_rate, mie_rate = rng.choice(TRIP_LOCATIONS)
        section = rng.choices(COMPANIES + ["Command Group", "Maintenance", "Medical"],
                              weights=[3, 3, 3, 3, 2, 2, 1])[0]
        doc = next_id("TA25-", 5)
        traveler = f"TRVLR {rng.randint(1000, 9999)}"
        posted = rand_date()
        nights = rng.randint(2, 12)
        base = f"TDY {purpose} {loc}"
        trip_rows = [
            add("21TA", f"{base} - AIRFARE", doc, posted, money(310, 920), "TDY TRAVEL", traveler, section),
            add("21TL", f"{base} - LODGING", doc, posted, nights * lodging_rate * rng.uniform(.95, 1.1),
                "TDY TRAVEL", traveler, section),
            add("21TM", f"{base} - M&IE", doc, posted, (nights + 0.5) * mie_rate, "TDY TRAVEL", traveler, section),
        ]
        if rng.random() < 0.65:
            trip_rows.append(add("21TC", f"{base} - RENTAL CAR", doc, posted,
                                 nights * money(38, 72), "TDY TRAVEL", traveler, section))
        if rng.random() < 0.45:
            trip_rows.append(add("21TX", f"{base} - MISC (BAG/PARKING/TOLLS)", doc, posted,
                                 money(35, 160), "TDY TRAVEL", traveler, section))

        roll = rng.random()
        if roll < 0.07:  # trip cancelled: everything de-obligated
            for r in trip_rows:
                add_adjustment(r, -r.amount, (3, 20))
        elif roll < 0.60:  # voucher settled for less than the estimate
            r = rng.choice(trip_rows[1:3])
            add_adjustment(r, -round(r.amount * rng.uniform(0.05, 0.35), 2), (nights + 5, nights + 35))


# ---------------------------------------------------------------------------
# Training: odd course names, seats bought per class.
# ---------------------------------------------------------------------------

COURSES = [
    ("CLS RECERT", 180, 450), ("MRT LVL I (MASTER RESILIENCE)", 600, 1400),
    ("SHARP SARC/VA 80HR", 900, 2100), ("COMBATIVES LVL 1 CERT", 250, 700),
    ("UPAR CRS", 150, 400), ("HAZMAT CERT DOT 49CFR", 600, 1500),
    ("FIELD SANITATION TEAM CRS", 120, 350), ("UNIT ARMORER CRS", 300, 800),
    ("CSDP TRNG - COMMAND SUPPLY DISCIPLINE", 200, 500), ("LSS GREEN BELT", 2200, 4800),
    ("ASIST SUICIDE INTERVENTION WKSHP", 250, 600), ("EO LEADER CRS", 400, 1100),
    ("CBRN DEFENSE CRS", 350, 900), ("BMSO LVL 2", 500, 1200),
    ("ACFT GRADER CERT", 90, 260), ("DRIVER TRNG TACT VEH (MDT)", 150, 450),
    ("TAMIS / RANGE OIC CERT", 100, 300), ("FORKLIFT 10K RT CERT", 450, 1100),
    ("HEAT INJURY PREVENTION TTT", 80, 220), ("ANTITERRORISM OFFICER LVL II", 600, 1600),
    ("GOV TRAVEL CARD APC TRNG", 150, 400), ("SPILL RESPONSE / ENV COMPLIANCE", 300, 800),
    ("COUNSELING & NCOER WRITING WKSHP", 120, 380), ("WILDERNESS FIRST RESPONDER", 750, 1500),
    ("LEADERS RECON & LAND NAV REFRESH", 90, 300), ("GENERATOR OPS & MAINT CRS", 500, 1300),
    ("CRANE OPERATOR CERT", 1500, 3200), ("WATER SURVIVAL / DUNKER TRNG", 400, 1000),
]


def make_training(count: int) -> None:
    class_numbers: dict[str, int] = {}
    for _ in range(count):
        course, low, high = rng.choice(COURSES)
        class_numbers[course] = class_numbers.get(course, 0) + 1
        seats = rng.randint(1, 6)
        section = rng.choice(COMPANIES + ["Medical", "Maintenance", "HQ Co"])
        desc = f"TUITION {course} CL 25-{class_numbers[course]:02d} ({seats} SEAT{'S' if seats > 1 else ''})"
        row = add("25TR", desc, next_id("TRN25-", 4), rand_date(), seats * money(low, high),
                  "INDIV TRAINING", next_id("REG", 6), section)
        if rng.random() < 0.18:  # a seat refunded / no-show credit
            per_seat = row.amount / seats
            add_adjustment(row, -round(per_seat * rng.randint(1, seats), 2))


# ---------------------------------------------------------------------------
# Government purchase card: recurring items bought many times.
# ---------------------------------------------------------------------------

GPC_ITEMS = [
    ("TONER HP 58X BLK", 140, 420, "HQ Co"), ("PAPER COPY 8.5X11 CASE", 45, 260, "HQ Co"),
    ("BAG PLASTIC 33GAL 250CT", 35, 120, "HQ Co"), ("ZIP TIES 14IN 1000PK", 25, 80, "Supply"),
    ("SANDBAGS 500CT", 180, 520, "Alpha Co"), ("PAINT OD GREEN 5GAL", 140, 380, "Maintenance"),
    ("WHITEBOARD MARKERS ASST", 15, 60, "HQ Co"), ("CABLE USB-C 6FT 10PK", 25, 70, "HQ Co"),
    ("BATTERIES AA 48PK", 20, 75, "Supply"), ("BATTERIES CR123A 24PK", 30, 95, "Alpha Co"),
    ("COFFEE MAKER COMMERCIAL", 180, 420, "Command Group"), ("FOLDING CHAIRS 10PK", 160, 380, "HQ Co"),
    ("LAMINATOR 12IN", 60, 180, "HQ Co"), ("PADLOCKS KEYED ALIKE 12PK", 70, 190, "Supply"),
    ("EXTENSION CORD 100FT HD", 45, 140, "Maintenance"), ("SHOP RAGS 50LB BOX", 55, 140, "Maintenance"),
    ("HAND SANITIZER 1GAL", 20, 60, "Medical"), ("NITRILE GLOVES XL 1000CT", 60, 180, "Maintenance"),
    ("CLEANING SUPPLIES - BAY", 60, 260, "Bravo Co"), ("WATER COOLER JUGS 5GAL", 30, 110, "Charlie Co"),
    ("ICE CHESTS 120QT", 120, 360, "Charlie Co"), ("POP-UP CANOPY 10X20", 220, 560, "Alpha Co"),
    ("TARPS 20X30 HD", 80, 240, "Bravo Co"), ("PARACORD 550 1000FT", 40, 110, "Alpha Co"),
    ("HEADLAMPS RED/WHITE 12PK", 90, 260, "Bravo Co"), ("CHEM LIGHTS GRN 100PK", 70, 180, "Charlie Co"),
    ("SHELVING UNIT STEEL 72IN", 160, 420, "Supply"), ("STORAGE TOTES 27GAL 8PK", 70, 180, "Supply"),
    ("LABEL MAKER TAPE 6PK", 35, 95, "Supply"), ("ENVELOPES 10X13 500CT", 30, 80, "HQ Co"),
    ("BINDERS 3IN 12PK", 35, 110, "HQ Co"), ("PRINTER DRUM UNIT", 120, 260, "HQ Co"),
    ("FLAGS / GUIDON REPLACEMENT", 80, 240, "Command Group"), ("AWARD PLAQUES 10CT", 180, 650, "Command Group"),
    ("COINS / FRAMES - RETIREMENT", 120, 480, "Command Group"), ("FIRST AID KITS OSHA 50P", 40, 130, "Medical"),
    ("EYEWASH STATION REFILL", 60, 160, "Maintenance"), ("FLOOR DRY ABSORBENT 40LB", 25, 75, "Maintenance"),
    ("DRIP PANS 24IN 6PK", 70, 170, "Maintenance"), ("AIR FILTERS 20X25 12PK", 60, 150, "HQ Co"),
    ("DEHUMIDIFIER 50PT", 220, 380, "Supply"), ("BOX FANS 20IN 4PK", 90, 210, "Charlie Co"),
    ("MOUSETRAPS / PEST CTRL", 25, 90, "HQ Co"), ("CONCRETE MIX 80LB 20BG", 110, 260, "Alpha Co"),
    ("LUMBER 2X4X8 PALLET", 380, 920, "Bravo Co"), ("SCREWS/HARDWARE ASST", 30, 160, "Maintenance"),
    ("HEARING PROTECTION 200PR", 70, 190, "Charlie Co"), ("SAFETY VESTS HI-VIS 20PK", 60, 150, "Maintenance"),
    ("TRAFFIC CONES 28IN 10PK", 120, 280, "Maintenance"), ("STRAPS RATCHET 2IN 8PK", 60, 160, "Supply"),
]


def make_gpc(count: int) -> None:
    for _ in range(count):
        item, low, high, section = rng.choice(GPC_ITEMS)
        cardholder = f"CH-{rng.choice(['0412', '0587', '0733', '0901'])}"
        row = add("26GP", f"GPC {item}", next_id("GPC25-", 5), rand_date(), money(low, high),
                  "GPC PURCHASES", cardholder, section)
        if rng.random() < 0.08:  # returned
            add_adjustment(row, -row.amount if rng.random() < 0.7 else -round(row.amount * 0.3, 2))


# ---------------------------------------------------------------------------
# Supply requisitions (NSN) and repair parts.
# ---------------------------------------------------------------------------

NSN_ITEMS = [
    ("6135-01-214-6072", "BATTERY,NONRECHARGEABLE BA-5590", 95, 140),
    ("6140-01-485-1472", "BATTERY,STORAGE 6TL", 280, 360),
    ("8465-01-580-0985", "PACK,FIELD MOLLE II", 140, 210),
    ("8415-01-628-9851", "GLOVES,COMBAT FLAME RESIST", 45, 70),
    ("8340-01-185-2602", "TENT,GENERAL PURPOSE SMALL", 900, 1500),
    ("7240-00-089-3827", "CAN,WATER MILITARY 5GAL", 18, 30),
    ("7240-01-337-5269", "CAN,FUEL MILITARY 5GAL", 25, 45),
    ("8970-00-149-1094", "MEAL,READY-TO-EAT CS", 110, 140),
    ("6545-01-532-3674", "KIT,FIRST AID INDIV IFAK", 55, 90),
    ("5120-01-578-0154", "TOOL KIT,GENERAL MECHANIC", 1400, 2600),
    ("4610-01-504-5698", "WATER PURIFICATION UNIT TABS", 40, 75),
    ("6230-01-517-1937", "FLASHLIGHT,TACTICAL", 35, 60),
    ("8430-01-514-5165", "BOOTS,COMBAT HOT WEATHER", 120, 165),
    ("8405-01-547-2559", "JACKET,COLD WEATHER GEN III", 160, 240),
    ("5985-01-454-9470", "ANTENNA,WHIP RT-1523", 220, 380),
    ("5820-01-451-8250", "HANDSET,RADIO H-250", 180, 260),
    ("7520-01-519-4372", "DESK,FIELD PORTABLE", 210, 340),
    ("4220-01-505-8964", "LIFE PRESERVER,VEST", 90, 140),
    ("6665-01-579-3437", "DETECTOR,CHEMICAL AGENT M8 PAPER", 18, 32),
    ("8105-00-857-2131", "BAG,SAND (500)", 120, 190),
]

PART_ITEMS = [
    "ALTERNATOR 200A M1083", "STARTER ASSY M1151", "TIRE 395/85R20 XZL", "TIRE 37X12.5R16.5 HMMWV",
    "BRAKE SHOE KIT FMTV", "FUEL FILTER/WATER SEP", "OIL ENGINE 15W40 55GAL", "COOLANT ANTIFREEZE 55GAL",
    "BATTERY CABLE ASSY", "HYDRAULIC HOSE ASSY LOADER", "WINDSHIELD MRAP", "SEAT BELT 5PT RESTRAINT",
    "CTIS VALVE ASSY", "INJECTOR FUEL 6.5L", "RADIATOR ASSY M1078", "LIGHT ASSY LED COMPOSITE",
    "GENERATOR VOLTAGE REG 10KW", "AIR COMPRESSOR GOVERNOR", "U-JOINT KIT DRIVESHAFT", "MIRROR ASSY SIDE",
    "WIRING HARNESS TRAILER", "PINTLE HOOK ASSY", "WHEEL BEARING KIT", "SERPENTINE BELT", "WATER PUMP ASSY",
    "TURBOCHARGER REMAN", "TRANSMISSION FILTER KIT", "DOOR HINGE KIT ARMORED", "FIRE EXTINGUISHER AFES BOTTLE",
    "SHOP EQUIP CALIBRATION SVC",
]


def make_requisitions(count: int) -> None:
    for _ in range(count):
        nsn, name, low, high = rng.choice(NSN_ITEMS)
        qty = rng.randint(1, 12)
        section = rng.choice(COMPANIES + ["Supply"])
        row = add("26SR", f"NSN {nsn} {name}", next_id("W25R", 6), rand_date(), qty * money(low, high),
                  "SUPPLY", f"QTY {qty}", section)
        if rng.random() < 0.13:  # cancelled requisition
            add_adjustment(row, -row.amount, (10, 60))


def make_parts(count: int) -> None:
    for _ in range(count):
        part = rng.choice(PART_ITEMS)
        section = rng.choices(["Maintenance", "Alpha Co", "Bravo Co", "Charlie Co"], weights=[6, 1, 1, 1])[0]
        row = add("26MX", f"PARTS {part}", next_id("WO25-", 5), rand_date(), money(85, 2900),
                  "MAINT", f"WO {rng.randint(100000, 999999)}", section)
        if rng.random() < 0.12:  # core credit / partial cancel
            add_adjustment(row, -round(row.amount * rng.uniform(0.1, 1.0), 2))


# ---------------------------------------------------------------------------
# Service contracts: one obligation, funded in monthly increments.
# ---------------------------------------------------------------------------

CONTRACTS = [
    ("PORTA JOHN SVC - TRNG AREAS FY25", 1400, 2600, "HQ Co"),
    ("COPIER LEASE & MAINT FY25", 900, 1300, "HQ Co"),
    ("LINEN & LAUNDRY SVC - BARRACKS", 1100, 1700, "HQ Co"),
    ("PEST CONTROL SVC BLDG 4410-4416", 350, 650, "HQ Co"),
    ("GENERATOR PREVENTIVE MAINT CONTRACT", 1600, 2800, "Maintenance"),
    ("VEH WASH RACK SVC", 600, 1200, "Maintenance"),
    ("HAZ WASTE PICKUP & DISPOSAL", 700, 1900, "Maintenance"),
    ("CELL PHONE SVC - DUTY PHONES", 800, 1100, "Command Group"),
    ("ICE MACHINE LEASE (4)", 300, 420, "Charlie Co"),
    ("RANGE BRASS & DEBRIS RECLAIM", 900, 2400, "Alpha Co"),
    ("FORKLIFT RENTAL 10K (2)", 1800, 2600, "Supply"),
    ("CONNEX STORAGE RENTAL (6)", 750, 900, "Supply"),
    ("TACTICAL VEH GPS TRACKING SUBSCRIPTION", 400, 650, "Maintenance"),
    ("MEDICAL WASTE DISPOSAL", 250, 500, "Medical"),
]


def make_contracts() -> None:
    for name, low, high, section in CONTRACTS:
        doc = next_id("PO-45001", 5)
        month = FY_START
        posted = FY_START + timedelta(days=rng.randint(0, 20))
        while posted <= FY_END:
            add("25SV", name, doc, posted, money(low, high), "BASE SVCS", "INCR FUNDING", section)
            month = (month.replace(day=28) + timedelta(days=5)).replace(day=1)
            posted = month + timedelta(days=rng.randint(10, 25))


# ---------------------------------------------------------------------------
# Fuel, medical, IT, and charges that belong to someone else.
# ---------------------------------------------------------------------------

def make_fuel() -> None:
    month = FY_START
    while month <= FY_END:
        tag = month.strftime("%b%y").upper()
        for fuel, low, high, section in [("JP8 BULK", 2200, 5200, "Maintenance"),
                                         ("DIESEL DF2 BULK", 1100, 2900, "Maintenance"),
                                         ("MOGAS UNLEADED", 300, 900, "HQ Co")]:
            add("26FL", f"FUEL {fuel} {tag}", f"FUEL25-{tag}", month + timedelta(days=rng.randint(2, 8)),
                money(low, high), "FUEL", "BULK ISSUE", section)
        month = (month.replace(day=28) + timedelta(days=5)).replace(day=1)


MED_ITEMS = [
    "CLS BAG RESTOCK", "AED PADS ADULT", "TOURNIQUETS CAT GEN7", "HEMOSTATIC GAUZE", "LITTER SLEDS",
    "BP CUFFS / STETHOSCOPES", "SPLINT SAM 36IN", "OXYGEN CYLINDER REFILL", "IV START KITS",
    "PULSE OXIMETERS", "EYE WASH BOTTLES", "HEAT INJURY ICE SHEETS", "SUNSCREEN / INSECT REPELLENT",
    "CHEST SEALS VENTED", "NASOPHARYNGEAL AIRWAYS", "TRAINING MANIKIN REPAIR",
]


def make_medical(count: int) -> None:
    for _ in range(count):
        item = rng.choice(MED_ITEMS)
        add("26MD", f"MED {item}", next_id("MED25-", 4), rand_date(), money(60, 1400),
            "MEDICAL", "PRIME VENDOR", "Medical")


IT_ITEMS = [
    ("25IT", "SOFTWARE LIC - MAINT MGMT ANNUAL", 3500, 6200, "Maintenance"),
    ("25IT", "SOFTWARE LIC - PDF EDITOR (25 SEATS)", 1200, 2200, "HQ Co"),
    ("25IT", "SOFTWARE LIC - SCHEDULING TOOL", 900, 1800, "Command Group"),
    ("31EQ", "LAPTOPS RUGGEDIZED (4)", 9800, 14500, "HQ Co"),
    ("31EQ", "VTC SYSTEM CONF RM", 6500, 11800, "Command Group"),
    ("31EQ", "PRINTER MFD COLOR A3", 3800, 6200, "HQ Co"),
    ("25IT", "MONITORS 27IN (10)", 1900, 2900, "HQ Co"),
    ("25IT", "CAC READERS / DOCKS", 400, 900, "HQ Co"),
    ("31EQ", "POWER GENERATOR 3KW PORTABLE (2)", 4200, 7400, "Alpha Co"),
    ("31EQ", "THERMAL IMAGING MONOCULARS (3)", 5200, 9800, "Bravo Co"),
]


def make_it() -> None:
    for code, name, low, high, section in IT_ITEMS:
        for _ in range(rng.randint(1, 2)):
            add(code, name, next_id("PO-45009", 5), rand_date(), money(low, high), "IT / EQUIPMENT",
                "CONTRACT", section)


MISCHARGES = [
    ("26MX", "PARTS ROTOR BLADE TIP CAP UH60", "WO25-A7731", 4875.00, "AVN MAINT"),
    ("25SV", "CATERING SVC BDE BALL 250 PAX", "PO-4500X1182", 8920.00, "BDE EVENTS"),
    ("21TL", "TDY BDE CDR CONF ARLINGTON VA - LODGING (ROOM BLOCK)", "TA25-X0091", 6240.00, "TDY TRAVEL"),
    ("26GP", "GPC MARINE OUTBOARD MOTOR SVC", "GPC25-X4410", 1385.40, "WATERCRAFT"),
    ("26SR", "NSN 1560-01-456-8812 PANEL,AIRCRAFT ACCESS", "W25RX00417", 2210.75, "AVN SUPPLY"),
    ("25TR", "TUITION HELO HOIST OPERATOR CRS CL 25-03 (2 SEATS)", "TRN25-X019", 3100.00, "AVN TRAINING"),
]


def make_mischarges() -> None:
    for code, desc, doc, amount, mission in MISCHARGES:
        add(code, desc, doc, rand_date(date(2025, 2, 1), date(2025, 8, 25)), amount, mission,
            "UNKNOWN", NOT_OURS, rng.choice(["FC-2040", "FC-3310", "FC-1188"]), exact=True)


# ---------------------------------------------------------------------------
# Writing the reports
# ---------------------------------------------------------------------------

def fiscal_month(d: date) -> int:
    return (d.month - 10) % 12 + 1


def write_report(path: Path, as_of: date, with_sections: bool) -> tuple[int, float]:
    wb = Workbook()
    ws = wb.active
    ws.title = "SOF"
    ws.append(HEADERS + (["Section"] if with_sections else []))
    for cell in ws[1]:
        cell.font = Font(bold=True)

    included = sorted((r for r in rows if r.posted <= as_of), key=lambda r: (r.posted, r.doc))
    total = 0.0
    for r in included:
        line = [
            "0100", r.funds_center, "OPS", r.code, fiscal_month(r.posted), r.desc, r.doc, r.posted,
            r.reference, f"WBS-{r.code}-25", r.mission, r.amount, r.amount,
        ]
        if with_sections:
            # The analyst's hand labels from last cycle; a few were never filled in.
            line.append(r.section if rng.random() > 0.02 else None)
        ws.append(line)
        total += r.amount

    for cell in ws["H"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for col in ("L", "M"):
        for cell in ws[col][1:]:
            cell.number_format = "#,##0.00"
    widths = {"A": 7, "B": 10, "C": 6, "D": 10, "E": 8, "F": 58, "G": 15, "H": 12,
              "I": 14, "J": 13, "K": 16, "L": 13, "M": 13, "N": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(included), total


def main() -> None:
    make_trips(78)
    make_training(95)
    make_gpc(400)
    make_requisitions(125)
    make_parts(140)
    make_contracts()
    make_fuel()
    make_medical(55)
    make_it()
    make_mischarges()

    out = Path(__file__).resolve().parent.parent / "sample_data"
    out.mkdir(exist_ok=True)
    n1, t1 = write_report(out / "SOF_period1.xlsx", PERIOD1_AS_OF, with_sections=True)
    n2, t2 = write_report(out / "SOF_period2.xlsx", PERIOD2_AS_OF, with_sections=False)
    full = sum(r.amount for r in rows)
    print(f"SOF_period1.xlsx  as of {PERIOD1_AS_OF}: {n1:,} rows, ${t1:,.2f} (labeled)")
    print(f"SOF_period2.xlsx  as of {PERIOD2_AS_OF}: {n2:,} rows, ${t2:,.2f}")
    print(f"Full fiscal year: {len(rows):,} rows, ${full:,.2f}")


if __name__ == "__main__":
    main()
